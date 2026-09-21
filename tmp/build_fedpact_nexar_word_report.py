from __future__ import annotations

import json
import math
import os
import random
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Python\実験_修正")
RUN_ROOT = ROOT / "artifacts" / "runs" / "fedpact_nexar_example"
BUILD_DIR = ROOT / "tmp" / "fedpact_nexar_word_report"
OUTPUT_DIR = ROOT / "output" / "documents"
OUTPUT_DOCX = OUTPUT_DIR / "FedPACT_NEXAR_Layer2_動作確認と局所更新診断_20260916.docx"

RISK_RUN = RUN_ROOT / "run_20260913_balanced_order101_risk_fullgrid"
NORMAL_RUN = RUN_ROOT / "run_20260913_balanced_order101_normal_fullgrid"
CONTROL_RUN = RUN_ROOT / "run_20260914_logit_scale_confirmation"

NAVY = "17365D"
BLUE = "2F6B9A"
LIGHT_BLUE = "EAF2F8"
PALE_BLUE = "F4F8FB"
ORANGE = "D97706"
RED = "B91C1C"
GREEN = "2F855A"
GRAY = "5B6573"
LIGHT_GRAY = "D9D9D9"
VERY_LIGHT_GRAY = "F5F6F7"
BLACK = "000000"
WHITE = "FFFFFF"


def load_json(path: Path):
    with path.open("r", encoding="utf-8") as f:
        return json.load(f)


JP_FONT = "Yu Gothic"
PIL_FONT_REGULAR = r"C:\Windows\Fonts\YuGothM.ttc"
PIL_FONT_BOLD = r"C:\Windows\Fonts\YuGothB.ttc"


def pil_font(size: int, bold: bool = False):
    return ImageFont.truetype(PIL_FONT_BOLD if bold else PIL_FONT_REGULAR, size=size)


def text_size(draw, text: str, font):
    box = draw.textbbox((0, 0), text, font=font)
    return box[2] - box[0], box[3] - box[1]


def draw_centered(draw, xy, text: str, font, fill="#111111"):
    w, h = text_size(draw, text, font)
    draw.text((xy[0] - w / 2, xy[1] - h / 2), text, font=font, fill=fill)


def save_chart(img: Image.Image, path: Path):
    img.save(path, format="PNG", dpi=(220, 220), optimize=True)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=LIGHT_GRAY, size=6):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        elem = borders.find(qn(tag))
        if elem is None:
            elem = OxmlElement(tag)
            borders.append(elem)
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), str(size))
        elem.set(qn("w:space"), "0")
        elem.set(qn("w:color"), color)


def set_cell_width(cell, inches: float):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def style_run(run, *, size=None, bold=None, color=BLACK, italic=None, font_name="Yu Gothic"):
    run.font.name = font_name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font_name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font_name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font_name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    return run


def style_paragraph_runs(paragraph, *, size=10.5, color=BLACK):
    for run in paragraph.runs:
        style_run(run, size=size, color=color)


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Yu Gothic"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Yu Gothic")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.18

    title = styles["Title"]
    title.font.name = "Yu Gothic"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Yu Gothic")
    title.font.size = Pt(25)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLACK)
    title.paragraph_format.space_after = Pt(14)
    title.paragraph_format.keep_with_next = True
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for name, size, before, after in [
        ("Heading 1", 16, 12, 7),
        ("Heading 2", 12.5, 10, 5),
        ("Heading 3", 11, 7, 3),
    ]:
        st = styles[name]
        st.font.name = "Yu Gothic"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Yu Gothic")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(BLACK)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True

    caption = styles["Caption"]
    caption.font.name = "Yu Gothic"
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "Yu Gothic")
    caption.font.size = Pt(9)
    caption.font.italic = False
    caption.font.color.rgb = RGBColor.from_string(GRAY)
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.keep_with_next = False


def add_title_block(doc: Document):
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.add_run("NEXAR実動画を用いたFedPACT第2層の動作確認と学習順の影響")

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("第I段階 Step I A および Step I B の一部")
    style_run(r, size=12, color=GRAY)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(16)
    r = p.add_run("吉村有生　2026年9月16日")
    style_run(r, size=10, color=GRAY)


def add_bullet(doc: Document, text: str, level: int = 0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.keep_with_next = False
    p.paragraph_format.left_indent = Inches(0.25 + 0.2 * level)
    p.paragraph_format.first_line_indent = Inches(-0.15)
    r = p.add_run(text)
    style_run(r, size=10.3)
    return p


def add_section_heading(doc: Document, text: str, level: int = 1, page_break: bool = True):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.page_break_before = page_break
    return p


def add_source_note(doc: Document, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run(text)
    style_run(r, size=8, color=GRAY)
    return p


def add_inline_notation(paragraph, pieces):
    """pieces: list of (text, italic, subscript, superscript)."""
    for text, italic, subscript, superscript in pieces:
        run = paragraph.add_run(text)
        style_run(run, size=10.5, italic=italic)
        run.font.subscript = subscript
        run.font.superscript = superscript


def set_alt_text_for_last_picture(paragraph, description: str):
    drawings = paragraph._p.xpath(".//w:drawing")
    if not drawings:
        return
    doc_prs = drawings[-1].xpath(".//wp:docPr")
    if doc_prs:
        doc_prs[0].set("descr", description)


def add_figure(doc: Document, image_path: Path, width: float, caption: str, alt_text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(image_path), width=Inches(width))
    set_alt_text_for_last_picture(p, alt_text)
    c = doc.add_paragraph(style="Caption")
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c.add_run(caption)
    style_run(r, size=9, color=GRAY)


def add_frame_strip(doc: Document, frame_paths, caption: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    for idx, frame_path in enumerate(frame_paths):
        if idx:
            spacer = p.add_run(" ")
            style_run(spacer, size=3)
        p.add_run().add_picture(str(frame_path), width=Inches(1.62))
    drawings = p._p.xpath(".//w:drawing")
    for idx, drawing in enumerate(drawings):
        doc_prs = drawing.xpath(".//wp:docPr")
        if doc_prs:
            doc_prs[0].set("descr", f"sample 8fab5a9e96 frame {idx + 1} of {len(drawings)}")
    c = doc.add_paragraph(style="Caption")
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c.add_run(caption)
    style_run(r, size=9, color=GRAY)


def add_table(doc: Document, headers, rows, widths, *, font_size=8.7, alignments=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for idx, (cell, text) in enumerate(zip(hdr.cells, headers)):
        set_cell_width(cell, widths[idx])
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(str(text))
        style_run(r, size=font_size, bold=True, color=WHITE)
    for row_idx, row_data in enumerate(rows):
        cells = table.add_row().cells
        for idx, (cell, value) in enumerate(zip(cells, row_data)):
            set_cell_width(cell, widths[idx])
            set_cell_shading(cell, WHITE if row_idx % 2 == 0 else PALE_BLUE)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            if alignments:
                p.alignment = alignments[idx]
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(value))
            style_run(r, size=font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def make_order_dependence_chart(path: Path):
    labels = ["学習順 A", "学習順 B", "学習順 C"]
    normal_pre = [0.222, 0.222, 0.222]
    normal_post = [0.121, 0.297, 0.080]
    risk_pre = [0.596, 0.596, 0.596]
    risk_post = [0.466, 0.707, 0.350]
    colors = [f"#{BLUE}", f"#{RED}", f"#{GREEN}"]
    img = Image.new("RGB", (1800, 1000), "white")
    d = ImageDraw.Draw(img)
    draw_centered(d, (900, 62), "学習する順番だけで出力変化の方向が反転", pil_font(44, True))
    draw_centered(d, (900, 122), "同じ学習用動画6本　同じ開始モデル　同じ学習条件　変えたのは学習順だけ", pil_font(27), "#5B6573")

    def draw_panel(x0, title, desired, pre_values, post_values):
        panel_w = 760
        d.rounded_rectangle((x0, 180, x0 + panel_w, 875), radius=20, outline="#CBD2DA", width=3)
        draw_centered(d, (x0 + panel_w / 2, 225), title, pil_font(34, True), "#162B49")
        draw_centered(d, (x0 + panel_w / 2, 270), "局所学習に使っていない動画3本の平均", pil_font(23), "#5B6573")
        d.text((x0 + 40, 315), desired, font=pil_font(23, True), fill="#7A4A00")
        axis_left, axis_right = x0 + 150, x0 + 710

        def xp(value):
            return axis_left + value * (axis_right - axis_left)

        for tick in [0.0, 0.2, 0.4, 0.6, 0.8, 1.0]:
            x = xp(tick)
            d.line((x, 360, x, 760), fill="#E3E7EB", width=2)
            draw_centered(d, (x, 795), f"{tick:.1f}", pil_font(20), "#4B5563")
        for idx, (label, pre, post, color) in enumerate(zip(labels, pre_values, post_values, colors)):
            y = 420 + idx * 145
            d.text((x0 + 28, y - 17), label, font=pil_font(25, True), fill="#23354D")
            x_pre, x_post = xp(pre), xp(post)
            d.line((x_pre, y, x_post, y), fill=color, width=8)
            direction = 1 if x_post >= x_pre else -1
            tip = x_post
            d.polygon([(tip, y), (tip - direction * 20, y - 13), (tip - direction * 20, y + 13)], fill=color)
            d.ellipse((x_pre - 11, y - 11, x_pre + 11, y + 11), fill="#9AA5B1", outline="#555555")
            d.ellipse((x_post - 12, y - 12, x_post + 12, y + 12), fill=color, outline="#333333")
            d.text((x_pre - 32, y - 52), f"{pre:.3f}", font=pil_font(20), fill="#555555")
            d.text((x_post - 32, y + 22), f"{post:.3f}", font=pil_font(20, True), fill=color)
        draw_centered(d, (x0 + panel_w / 2, 842), "平均 q(risk)", pil_font(24), "#23354D")

    draw_panel(80, "normal動画　危険ではない", "望ましい方向　← risk確率が下がる", normal_pre, normal_post)
    draw_panel(960, "risk動画　危険である", "望ましい方向　risk確率が上がる →", risk_pre, risk_post)
    draw_centered(d, (900, 945), "灰色が局所学習前　色付きが局所学習後", pil_font(24), "#5B6573")
    save_chart(img, path)


def make_transformation_examples(trace, path: Path):
    candidates = trace["transformation_and_optimization"]["candidates"]
    selected = trace["quality_and_selection"]["selected"]
    one_seed = [c for c in candidates if c["seed_id"] == selected["seed_id"]]
    img = Image.new("RGB", (1800, 980), "white")
    d = ImageDraw.Draw(img)
    draw_centered(d, (900, 50), "server側の元動画1本から作成した10通りの変換候補", pil_font(40, True))
    draw_centered(d, (900, 103), "cropは2通り　ほかに明度・コントラスト・ぼかしを含む", pil_font(25), "#5B6573")

    labels = [
        "変換なし", "明度 0.75", "明度 0.90", "明度 1.10", "明度 1.25",
        "コントラスト 0.80", "コントラスト 1.20", "中央crop 0.85", "中央crop 0.92", "ぼかし 0.75",
    ]
    for idx, (candidate, label) in enumerate(zip(one_seed, labels)):
        row, col = divmod(idx, 5)
        x = 55 + col * 345
        y = 165 + row * 340
        frame = Image.open(RISK_RUN / "server" / candidate["frame_files"][0]).convert("RGB")
        frame.thumbnail((310, 190))
        canvas = Image.new("RGB", (310, 190), "white")
        canvas.paste(frame, ((310 - frame.width) // 2, (190 - frame.height) // 2))
        img.paste(canvas, (x, y))
        outline = "#D97706" if idx == 7 else "#B9C2CC"
        width = 7 if idx == 7 else 3
        d.rectangle((x, y, x + 310, y + 190), outline=outline, width=width)
        draw_centered(d, (x + 155, y + 228), label, pil_font(22, True), "#7A4A00" if idx == 7 else "#23354D")

    d.line((85, 865, 1715, 865), fill="#D8DEE6", width=3)
    draw_centered(d, (900, 915), "この例では中央crop 0.85の L match が最小", pil_font(29, True), "#7A4A00")
    save_chart(img, path)


def make_layer2_overview(path: Path):
    img = Image.open(RISK_RUN / "presentation_example.png").convert("RGB")
    d = ImageDraw.Draw(img)
    blue_fill = "#EAF2F8"
    background = "#F4F8FB"

    d.rounded_rectangle((45, 470, 1755, 750), radius=24, fill=blue_fill, outline="#7890A5", width=2)
    d.text((70, 492), "SERVER SEED（変換前）", font=pil_font(31, True), fill="#162B49")
    selected = load_json(RISK_RUN / "server" / "layer2_trace.json")["quality_and_selection"]["selected"]
    frame = Image.open(RISK_RUN / "server" / selected["frame_files"][0]).convert("RGB")
    frame.thumbnail((260, 145))
    for x in (70, 330, 590, 850):
        img.paste(frame, (x, 555))
    d.text((1150, 560), "8本のserver seedから候補を探索", font=pil_font(25, True), fill="#183F73")
    d.text((1150, 610), "各候補に10通りの変換を適用", font=pil_font(25), fill="#183F73")
    d.text((1150, 660), "この例では中央crop 0.85を選択", font=pil_font(25), fill="#183F73")

    d.rectangle((45, 755, 1755, 790), fill=background)
    d.text((70, 758), "↓ 候補探索 → 10通りの変換 → 目的関数による評価と選択", font=pil_font(23), fill="#315887")
    d.rectangle((45, 1120, 1755, 1178), fill=background)
    d.text((60, 1138), "※ 上段のprivate動画は説明用のローカル表示であり、clientからserverへ送信しない。", font=pil_font(20), fill="#9A3D12")
    save_chart(img, path)


def make_diagnostic_bars(path: Path):
    conditions = [
        "LoRA+head\n勾配累積",
        "LoRA+head\n層化batch",
        "headのみ",
        "LoRAのみ",
        "LoRAのみ\n3 epoch",
        "LoRA+head\n学習12本 *",
        "誤差重み\n学習6本",
        "誤差重み\n学習12本 *",
        "logit倍率\n対照",
    ]
    normal = [-0.0199, 0.0287, -0.0160, -0.0041, -0.0215, 0.0791, 0.0062, 0.0145, -0.1548]
    risk = [-0.0218, 0.0825, -0.0223, 0.0004, -0.0016, 0.0844, 0.0184, 0.0222, 0.0803]
    img = Image.new("RGB", (1800, 1250), "white")
    d = ImageDraw.Draw(img)
    draw_centered(d, (900, 55), "通常の局所更新では目的方向が評価用動画で安定しない", pil_font(40, True))
    left_label, plot_left, plot_right = 30, 690, 1720
    top, row_h = 150, 102
    xmin, xmax = -0.18, 0.10

    def xp(v):
        return plot_left + (v - xmin) / (xmax - xmin) * (plot_right - plot_left)

    zero = xp(0)
    for tick in [-0.15, -0.10, -0.05, 0, 0.05, 0.10]:
        x = xp(tick)
        d.line((x, top - 15, x, top + row_h * len(conditions)), fill="#E2E6EA", width=2)
        draw_centered(d, (x, top + row_h * len(conditions) + 35), f"{tick:+.2f}", pil_font(22))
    d.line((zero, top - 20, zero, top + row_h * len(conditions)), fill="#333333", width=3)
    for i, label in enumerate(conditions):
        y_mid = top + i * row_h + row_h / 2
        if i in (5, 7):
            d.rectangle((left_label - 10, y_mid - row_h / 2 + 3, plot_right, y_mid + row_h / 2 - 3), fill="#F3F4F6")
            for tick in [-0.15, -0.10, -0.05, 0, 0.05, 0.10]:
                x = xp(tick)
                d.line((x, y_mid - row_h / 2 + 3, x, y_mid + row_h / 2 - 3), fill="#E2E6EA", width=2)
            d.line((zero, y_mid - row_h / 2 + 3, zero, y_mid + row_h / 2 - 3), fill="#333333", width=3)
        display = label.replace("\n", " ")
        d.text((left_label, y_mid - 20), display, font=pil_font(23), fill="#23354D")
        for offset, value, color in [(-17, normal[i], BLUE), (17, risk[i], RED)]:
            y = y_mid + offset
            xval = xp(value)
            x0, x1 = sorted((zero, xval))
            d.rectangle((x0, y - 10, x1, y + 10), fill=f"#{color}")
            tx = xval + (8 if value >= 0 else -88)
            d.text((tx, y - 17), f"{value:+.3f}", font=pil_font(20), fill="#222222")
    d.rectangle((710, 105, 730, 125), fill=f"#{BLUE}")
    d.text((740, 98), "評価用 normal", font=pil_font(22), fill="#333333")
    d.rectangle((1020, 105, 1040, 125), fill=f"#{RED}")
    d.text((1050, 98), "評価用 risk", font=pil_font(22), fill="#333333")
    d.text((690, 1160), "目的方向　normal < -0.01　risk > +0.01", font=pil_font(24), fill="#5B6573")
    d.text((690, 1200), "* 学習12本の条件は評価対象が異なり、学習6本の条件と直接比較しない", font=pil_font(21), fill="#5B6573")
    save_chart(img, path)


def make_control_metrics(control, path: Path):
    groups = [("学習用", "train"), ("評価用", "evaluation"), ("確認用", "confirmation")]
    metrics = [("accuracy", "Accuracy", (0, 1.08)), ("ce", "Cross entropy", (0, 0.56)), ("brier", "Brier score", (0, 0.18))]
    img = Image.new("RGB", (1800, 900), "white")
    d = ImageDraw.Draw(img)
    draw_centered(d, (900, 55), "logit倍率対照では予測ラベルを変えず確信度指標だけが改善", pil_font(39, True))
    panel_lefts = [80, 625, 1170]
    panel_w, plot_top, plot_bottom = 500, 190, 720
    for panel_left, (key, title, ylim) in zip(panel_lefts, metrics):
        draw_centered(d, (panel_left + panel_w / 2, 140), title, pil_font(30, True), "#23354D")
        axis_left, axis_right = panel_left + 55, panel_left + panel_w - 15
        d.line((axis_left, plot_top, axis_left, plot_bottom), fill="#333333", width=2)
        d.line((axis_left, plot_bottom, axis_right, plot_bottom), fill="#333333", width=2)

        def yp(v):
            return plot_bottom - (v - ylim[0]) / (ylim[1] - ylim[0]) * (plot_bottom - plot_top)

        for frac in [0, 0.25, 0.5, 0.75, 1.0]:
            val = ylim[0] + (ylim[1] - ylim[0]) * frac
            y = yp(val)
            d.line((axis_left, y, axis_right, y), fill="#E5E8EB", width=1)
        group_centers = [axis_left + 85, axis_left + 220, axis_left + 355]
        for center, (label, g) in zip(group_centers, groups):
            pre = control[g]["pre"][key]
            post = control[g]["post"][key]
            for x, value, color in [(center - 25, pre, "AAB7C4"), (center + 25, post, GREEN)]:
                y = yp(value)
                d.rectangle((x - 20, y, x + 20, plot_bottom), fill=f"#{color}")
                draw_centered(d, (x, y - 20), f"{value:.3f}", pil_font(18), "#333333")
            draw_centered(d, (center, plot_bottom + 35), label, pil_font(21))
    d.rectangle((700, 800, 725, 825), fill="#AAB7C4")
    d.text((738, 792), "適応前", font=pil_font(23), fill="#333333")
    d.rectangle((900, 800, 925, 825), fill=f"#{GREEN}")
    d.text((938, 792), "適応後", font=pil_font(23), fill="#333333")
    save_chart(img, path)


def build_report():
    BUILD_DIR.mkdir(parents=True, exist_ok=True)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    layer2_trace = load_json(RISK_RUN / "server" / "layer2_trace.json")

    order_chart = BUILD_DIR / "order_dependence_chart.png"
    transformation_examples = BUILD_DIR / "transformation_examples.png"
    layer2_overview = BUILD_DIR / "layer2_overview_clean.png"
    make_order_dependence_chart(order_chart)
    make_transformation_examples(layer2_trace, transformation_examples)
    make_layer2_overview(layer2_overview)

    doc = Document()
    configure_document(doc)
    doc.core_properties.title = "NEXAR実動画を用いたFedPACT第2層の動作確認と学習順の影響"
    doc.core_properties.subject = "FedPACT 第I段階 Step I A および Step I B の一部に関する実験報告"
    doc.core_properties.author = "吉村有生"
    doc.core_properties.keywords = "FedPACT, NEXAR, Layer 2, 局所更新, proxy"

    add_title_block(doc)

    doc.add_heading("要約", level=1)
    p = doc.add_paragraph()
    r = p.add_run("結論　")
    style_run(r, bold=True)
    r = p.add_run("1 client・1 roundで、局所学習から代理データとlogit prototypeのpair作成までを一続きに確認した。一方、局所学習後の出力は動画の学習順によって変わった。また、riskのlogit prototypeに対して、normalラベルの元動画を変換した候補が選ばれた。出力確率の近さだけでなく、元動画のclass labelや内容も確認する必要がある。")
    style_run(r)

    add_bullet(doc, "確認済み　局所適応、クラスごとの平均確率によるlogit prototype、server seed検索、変換候補の評価、代理ペア生成")
    add_bullet(doc, "主要な問題　局所学習後の出力が学習順で変わること、L matchが最小の候補と目的classのlabelが一致しない場合があること")
    add_bullet(doc, "未実施　複数clientの連合集約、第3層のfoundation LoRA更新、3から5 round接続")

    p = doc.add_paragraph()
    r = p.add_run("本報告の位置づけ　")
    style_run(r, bold=True)
    r = p.add_run("FedPACT実験計画の第I段階のうち、clientでの適応出力取得、logit prototype作成、server側シードデータの探索・変換、代理データとlogit prototypeのpair作成を確認した途中報告である。性能比較や知識転移の結果ではない。")
    style_run(r)

    add_section_heading(doc, "1 実験目的と確認範囲")
    p = doc.add_paragraph()
    r = p.add_run("NEXARのnormalとriskの実動画を用い、client内の局所適応で得た出力変化をprototypeへ集約し、server側のseed検索と変換を経て、第3層へ渡すpairを生成できるかを確認した。")
    style_run(r)

    p = doc.add_paragraph()
    p.add_run("確認した経路　")
    style_run(p.runs[-1], bold=True)
    add_inline_notation(p, [
        ("q", True, False, False), ("pre", False, False, True),
        (" と ", False, False, False),
        ("q", True, False, False), ("post", False, False, True),
        (" から Δq を計算し、", False, False, False),
        ("μ", True, False, False), ("post", False, False, True), ("k,r", False, True, False),
        (" と Δ", False, False, False), ("μ", True, False, False), ("k,r", False, True, False),
        (" を作成する。その後、server seedを検索・変換し、", False, False, False),
        ("P", True, False, False), ("k,r", False, True, False),
        (" = (", False, False, False), ("x̃", True, False, False), ("k,r", False, True, False),
        (", ", False, False, False), ("μ", True, False, False), ("post", False, False, True), ("k,r", False, True, False),
        (") を生成する。", False, False, False),
    ])

    p = doc.add_paragraph()
    add_inline_notation(p, [
        ("q", True, False, False), ("pre", False, False, True),
        (" と ", False, False, False),
        ("q", True, False, False), ("post", False, False, True),
        (" は各動画に対する局所学習前後のsoftmax出力確率、Δqはその差である。", False, False, False),
        ("μ", True, False, False), ("post", False, False, True), ("k,r", False, True, False),
        (" はclass rのprivate動画について局所学習後の出力確率を平均したlogit prototype、Δ", False, False, False),
        ("μ", True, False, False), ("k,r", False, True, False),
        (" は局所学習前後のclass平均の差である。", False, False, False),
    ])

    doc.add_heading("実験条件", level=2)
    add_table(
        doc,
        ["項目", "設定"],
        [
            ["モデル", "Qwen2.5 VL 3B Instruct"],
            ["client", "1 client 1 round"],
            ["局所学習用動画", "normal 3本 risk 3本"],
            ["評価用動画", "normal 3本 risk 3本　局所学習・prototype作成・server検索に未使用"],
            ["server seed", "normal 4本 risk 4本"],
            ["prototype", "クラスごとの平均確率によるlogit prototype"],
            ["出力表現", "qとμはsoftmax出力確率　Δq = q post - q pre"],
            ["frequency", "π k,r = 3　各logit prototypeを作成した動画数"],
            ["候補探索", "8 seed × 10変換 = 80候補"],
            ["目的関数", "post KL + delta MSE　各重み1.0"],
            ["未導入", "L reg Qinv Layer 3"],
        ],
        [1.7, 5.2],
        font_size=9.2,
        alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
    )
    add_source_note(doc, "対象　2026年9月13日に実施したrisk及びnormalの第2層動作確認")

    p = doc.add_paragraph()
    r = p.add_run("標本数と評価範囲　")
    style_run(r, bold=True)
    r = p.add_run("各class n=3、反復なしの小規模な動作確認である。これらの評価用動画は局所学習、prototype作成、server seed検索には使用していない。ただし、後続の局所学習条件を選ぶためにも使用した。そのため、最終的な性能評価に使う、条件選択にも未使用のtest動画ではない。また、risk動画はevent時点、normal動画はrandom startからframeを抽出しており、frame sampling方法とclassが交絡している。")
    style_run(r)

    p = doc.add_paragraph()
    r = p.add_run("1 clientであることの制約　")
    style_run(r, bold=True)
    r = p.add_run("今回は1 clientだけであるため、連合集約後のglobal surrogate ")
    style_run(r)
    add_inline_notation(p, [
        ("M", True, False, False), ("G", False, True, False), ("t+1", False, False, True),
        (" は、そのclientの局所学習後surrogateと同じになる。複数clientの連合集約によって知識が保持されるか、薄まるかは確認していない。", False, False, False),
    ])

    doc.add_heading("clientからserverへ送信する適応情報", level=2)
    p = doc.add_paragraph()
    r = p.add_run("今回clientからserverへ送信する適応情報には、LoRA更新情報、局所学習後のprototype、局所変化量、frequencyを含めた。private動画、raw frame、sample ID、局所学習前prototypeの独立送信は含めていない。ただし、これは送信内容の構成を確認した結果であり、通信経路とserver側filesystem全体を対象としたprivacy保証ではない。")
    style_run(r)

    add_section_heading(doc, "2 NEXAR実動画による第2層の具体例", page_break=False)
    p = doc.add_paragraph()
    r = p.add_run("riskのlogit prototypeを具体例として追跡した。risk動画3本の平均risk確率は局所学習前0.685、局所学習後0.596となり、平均で0.090低下した。したがって、この例で第2層が再現する対象は、有用性が確認されたrisk知識ではなく、実際に観測された局所適応挙動である。")
    style_run(r)

    add_figure(
        doc,
        layer2_overview,
        6.95,
        "図1　NEXAR実動画から代理データ・logit prototypeのpairを生成するまでの具体例",
        "NEXARのclient動画、server seed、変換後proxyを上から順に示す第2層の具体例",
    )

    add_table(
        doc,
        ["対象", "局所学習後の平均出力", "局所変化量", "選択した変換", "post KL", "delta MSE", "L match"],
        [
            ["risk", "[0.404, 0.596]", "[+0.090, -0.090]", "中央crop 0.85", "0.000109", "0.000169", "0.000279"],
            ["normal", "[0.799, 0.201]", "[+0.139, -0.139]", "明度 0.75", "0.000080", "0.000389", "0.000468"],
        ],
        [0.6, 1.45, 1.4, 1.25, 0.75, 0.8, 0.75],
        font_size=8.2,
    )
    add_source_note(doc, "数値は各runのsummary.jsonとserver layer2 traceから取得。成分順は normal risk。")

    add_section_heading(doc, "3 10通りの変換候補と代理データの評価", page_break=False)
    p = doc.add_paragraph()
    r = p.add_run("8本のserver seedそれぞれに10通りの変換候補を適用し、合計80候補のL matchを比較した。10通りはcropだけではなく、変換なし1通り、明度4通り、コントラスト2通り、中央crop 2通り、ガウシアンぼかし1通りである。例えば中央crop 0.85は、各frameの中央85%を残して元の大きさへ拡大する処理であり、動画から取り出した4枚すべてに同じ処理を適用した。")
    style_run(r)

    add_figure(
        doc,
        transformation_examples,
        6.8,
        "図2　server側の元動画1本から作成した10通りの変換候補",
        "変換なし、明度4通り、コントラスト2通り、中央crop 2通り、ぼかし1通りの実画像例",
    )

    add_table(
        doc,
        ["変換の種類", "設定値", "L match"],
        [
            ["変換なし", "-", "0.062256"],
            ["明度", "0.75", "0.089976"],
            ["明度", "0.90", "0.038402"],
            ["明度", "1.10", "0.031404"],
            ["明度", "1.25", "0.028073"],
            ["コントラスト", "0.80", "0.074969"],
            ["コントラスト", "1.20", "0.052749"],
            ["中央crop", "0.85", "0.000279"],
            ["中央crop", "0.92", "0.002416"],
            ["ガウシアンぼかし", "0.75", "0.040767"],
        ],
        [3.1, 1.6, 2.2],
        font_size=8.8,
    )

    p = doc.add_paragraph()
    r = p.add_run("解釈　")
    style_run(r, bold=True)
    r = p.add_run("この例では中央crop 0.85のL matchが0.000279で最小となった。ただし、riskのlogit prototypeに対して選ばれた元動画のclass labelはnormalであった。動画内容の妥当性は人手で独立に評価していないため、内容が不一致とは断定しない。今回の候補選択はL matchだけを用い、L regとQinvは未導入である。また、候補選択と評価に同じL matchを用いているため、最小値は独立した品質評価ではない。第3層へ渡す前に、class labelとの整合と動画内容の確認が必要である。")
    style_run(r)

    doc.add_heading("この結果から確認できたこと", level=2)
    add_bullet(doc, "8本のserver seedに10通りずつの変換を適用し、80候補を比較できた")
    add_bullet(doc, "中央cropの比率を変えるだけでもL matchは大きく変化した")
    add_bullet(doc, "出力確率の近さと、元動画のclass label・内容の妥当性を分けて確認する必要がある")

    add_section_heading(doc, "4 局所更新の学習順依存", page_break=False)
    p = doc.add_paragraph()
    r = p.add_run("同じ6本の動画、同じ開始モデル、同じ局所学習条件を用い、6本を学習させる順番だけを変えた3条件を学習順A・B・Cとした。図3は、局所学習に使っていない動画に対する平均risk確率が、局所学習の前後でどう変化したかを示す。")
    style_run(r)

    add_figure(
        doc,
        order_chart,
        6.0,
        "図3　学習順だけを変えたときの平均risk確率の変化",
        "局所学習に使っていないnormal動画とrisk動画について、学習順A・B・Cごとに局所学習前後の平均risk確率を矢印で示すグラフ",
    )

    p = doc.add_paragraph()
    r = p.add_run("読み方　")
    style_run(r, bold=True)
    r = p.add_run("灰色の点は、3条件で共通する局所学習前の値である。色付きの点は局所学習後の値であり、矢印の向きが変化方向を表す。normal動画では左向き、risk動画では右向きが望ましい。学習順AとCでは両方の動画群でrisk確率が下がり、学習順Bでは両方で上がった。どの学習順もnormalとriskを同時に望ましい方向へ動かしていない。")
    style_run(r)

    p = doc.add_paragraph()
    r = p.add_run("判断　")
    style_run(r, bold=True)
    r = p.add_run("同じ動画と条件でも、学習順だけで局所学習後の出力方向が変わった。ただし、各学習順は1回ずつしか実施していないため、学習順依存の大きさや再現性を結論づけるには反復が必要である。")
    style_run(r)

    add_section_heading(doc, "5 結論と次の検証", page_break=False)
    doc.add_heading("観測事実から言えること", level=2)
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.first_line_indent = Inches(-0.15)
    add_inline_notation(p, [
        ("1 clientで、局所学習から ", False, False, False),
        ("P", True, False, False), ("k,r", False, True, False),
        (" = (", False, False, False), ("x̃", True, False, False), ("k,r", False, True, False),
        (", ", False, False, False), ("μ", True, False, False), ("post", False, False, True), ("k,r", False, True, False),
        (") の生成までを一続きに動作確認できた", False, False, False),
    ])
    add_bullet(doc, "riskとnormalの双方について、代理データとlogit prototypeのpairを生成できた")
    add_bullet(doc, "同じ動画と局所学習条件でも、学習順によって局所学習後の出力方向が変わった")
    add_bullet(doc, "riskのlogit prototypeに対し、L matchが最小の候補はnormalラベルの元動画を変換したものだった")
    add_bullet(doc, "frame sampling方法がclassと交絡しているため、観測差を一般的なrisk認識性能として扱えない")

    doc.add_heading("現時点では言えないこと", level=2)
    add_bullet(doc, "normal risk識別性能が改善した")
    add_bullet(doc, "未見動画へ一般化した")
    add_bullet(doc, "proxyが意味的なclient知識を再現し、foundation modelへ知識が転移した")
    add_bullet(doc, "複数clientの連合集約と3から5 roundの状態継承が正常に動作した")

    doc.add_heading("次の優先実験", level=2)
    steps = [
        ("1", "学習順の反復確認", "動画ごとの学習回数とoptimizer step数をそろえ、複数の学習順と乱数条件で同じ傾向が再現するか確認する。"),
        ("2", "2 client・1 roundの第1層確認", "各clientの局所変化量と、連合集約前後のglobal surrogate出力差を分けて保存する。"),
        ("3", "第2層の候補内容確認", "class labelが一致する候補、ランダム候補、検索候補を比較し、L matchだけで候補を決めない。"),
        ("4", "第3層と複数round接続", "foundation LoRA更新を1 roundで確認した後、3 round、状態継承が正常なら5 roundへ進む。"),
    ]
    add_table(doc, ["順序", "作業", "確認内容"], steps, [0.6, 1.85, 4.45], font_size=8.7,
              alignments=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT])

    add_section_heading(doc, "6 参照記録", page_break=False)
    sources = [
        "FedPACT実験計画260901.pdf　第I段階 Step I A・Step I B",
        "FedPACT手法提案260901.pdf　第2層",
        "2026年9月13日 NEXAR Layer 2 risk・normal実験ログ",
        "2026年9月13日 学習順A・B・C診断ログ",
    ]
    for source in sources:
        add_bullet(doc, source)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    r = p.add_run("記録上の注意　")
    style_run(r, bold=True, size=9.5)
    r = p.add_run("9月13日の実験は第I段階の小規模な動作確認である。各class n=3、各学習順1回のみであり、数値は性能推定ではなく、処理経路と学習順による変化の確認に使用する。")
    style_run(r, size=9.5)

    doc.save(OUTPUT_DOCX)
    print(str(OUTPUT_DOCX))
    print(f"FONT={JP_FONT}")


if __name__ == "__main__":
    build_report()
